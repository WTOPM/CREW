"""Trim junk from UN list and align Data VLOOKUP ranges after Chapter 3.2 refresh."""
from __future__ import annotations

import re
from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula

SRC = r"C:\Users\wtopm\Downloads\IMDG+IMS\IMDG_list_optimized.xlsx"


def main() -> None:
    wb = load_workbook(SRC)
    ws = wb["UN list"]

    last = 1
    bad_rows: list[int] = []
    for r in range(2, (ws.max_row or 1) + 1):
        v = ws.cell(r, 1).value
        if v is None:
            continue
        s = str(v).strip()
        if s.startswith("="):
            bad_rows.append(r)
            continue
        try:
            int(float(s))
            last = r
        except Exception:
            bad_rows.append(r)

    print("last data row", last, "bad", bad_rows[:20], "count bad", len(bad_rows))

    # Clear formula junk rows in-place first (within data block)
    for r in bad_rows:
        if r <= last:
            for c in range(1, 8):
                ws.cell(r, c).value = None

    # Recompute last after clearing mid-block junk
    last = 1
    for r in range(2, (ws.max_row or 1) + 1):
        v = ws.cell(r, 1).value
        if v is None:
            continue
        s = str(v).strip()
        if s.startswith("="):
            continue
        try:
            int(float(s))
            last = r
        except Exception:
            continue

    if (ws.max_row or 0) > last:
        ws.delete_rows(last + 1, ws.max_row - last)
        print("deleted rows after", last)

    for r in range(1, last + 1):
        for c in range(8, 50):
            if ws.cell(r, c).value is not None:
                ws.cell(r, c).value = None

    print("new max_row", ws.max_row, "max_col", ws.max_column)

    data = wb["Data"]
    range_pg = f"'UN list'!$D$2:$D${last}"
    range_tbl = f"'UN list'!$A$2:$G${last}"
    updated = 0
    for row in data.iter_rows(min_row=4, max_row=data.max_row, min_col=13, max_col=17):
        for cell in row:
            v = cell.value
            text = None
            if isinstance(v, ArrayFormula):
                text = v.text
            elif isinstance(v, str) and "UN list" in v:
                text = v
            if not text or "UN list" not in text:
                continue
            t = text if text.startswith("=") else f"={text}"
            t = re.sub(r"'UN list'!\$D\$2:\$D\$\d+", range_pg, t)
            t = re.sub(r"'UN list'!\$A\$2:\$G\$\d+", range_tbl, t)
            cell.value = ArrayFormula(cell.coordinate, t)
            updated += 1

    print("formulas updated", updated, "range last", last)
    wb.save(SRC)
    print("saved", SRC)

    wb2 = load_workbook(SRC)
    ws2 = wb2["UN list"]
    print(
        "verify max_row",
        ws2.max_row,
        "A last",
        ws2.cell(ws2.max_row, 1).value,
        "desc",
        str(ws2.cell(ws2.max_row, 2).value)[:60],
    )
    count = sum(1 for r in range(2, ws2.max_row + 1) if ws2.cell(r, 1).value is not None)
    print("count", count)
    c4 = wb2["Data"].cell(4, 13).value
    print("formula sample", getattr(c4, "text", c4))


if __name__ == "__main__":
    main()
